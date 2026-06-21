# -*- coding: utf-8 -*-
"""
Build report/SRS_ManageTicketTracking.docx.

Reuses report/SRS_ManageBacklog.docx as a *style template* (same cover layout,
Aptos Display headings, table look) so the new SRS matches its format and
composition exactly. The Ticket-Tracking screenshots are inserted full-width
(width only -> aspect ratio preserved, never cropped).
"""
import os, zipfile
from docx import Document
from docx.shared import Pt, Emu, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT      = os.path.dirname(os.path.abspath(__file__))
TEMPLATE  = os.path.join(ROOT, 'SRS_ManageBacklog.docx')
OUT       = os.path.join(ROOT, 'SRS_ManageTicketTracking.docx')
IMG_DIR   = os.path.normpath(os.path.join(ROOT, '..', 'force-app', 'Ticket-Tracking-images'))
TMP_LOGO  = os.path.join(ROOT, '_cover_logo.png')

# ── extract the cover logo (word/media/image1.png) from the template ──────────
with zipfile.ZipFile(TEMPLATE) as z:
    with open(TMP_LOGO, 'wb') as f:
        f.write(z.read('word/media/image1.png'))

doc = Document(TEMPLATE)

# ── wipe the template body, keep the trailing sectPr (page setup) ─────────────
body = doc.element.body
for child in list(body):
    if child.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(child)

CONTENT_W = Inches(6.3)            # fits inside 1" margins on Letter
HEADING_COLOR = RGBColor(0x0F, 0x47, 0x61)

# ── helpers ──────────────────────────────────────────────────────────────────
def cover_line(text, size, italic=True, black=False, before=None, after=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if before is not None: p.paragraph_format.space_before = Emu(before)
    if after  is not None: p.paragraph_format.space_after  = Emu(after)
    r = p.add_run(text)
    r.italic = italic
    r.font.name = 'Aptos Display'
    r.font.size = Pt(size)
    if black:
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def h(level, text):
    return doc.add_heading(text, level=level)

def normal(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def bullet(text):
    return doc.add_paragraph('•  ' + text)

def figure(filename, caption):
    path = os.path.join(IMG_DIR, filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=CONTENT_W)   # width only -> no crop
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(0x6B, 0x77, 0x8C)

def set_cell_bold(cell):
    for para in cell.paragraphs:
        for r in para.runs:
            r.bold = True

def actors_table(rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    # widths ~ 4680 dxa each (matches template's two equal columns)
    for r_idx, (a, b) in enumerate(rows):
        c0, c1 = table.rows[r_idx].cells
        c0.text, c1.text = a, b
        c0.width = Emu(2971800); c1.width = Emu(2971800)
    set_cell_bold(table.rows[0].cells[0]); set_cell_bold(table.rows[0].cells[1])
    # borders: bottom + insideH single sz4 D0D7DE  (template look)
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('bottom', 'insideH'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0');    el.set(qn('w:color'), 'D0D7DE')
        borders.append(el)
    tblPr.append(borders)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), '9360'); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)
    return table

# ═════════════════════════════════════════════════════════════════════════════
# COVER
# ═════════════════════════════════════════════════════════════════════════════
img_p = doc.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.paragraph_format.space_before = Emu(1524000)
img_p.add_run().add_picture(TMP_LOGO, width=Inches(3.12))

cover_line('Technical Document', 24, black=True, before=381000, after=76200)
cover_line('Manage Ticket Tracking — SRS', 18, black=True, after=254000)
cover_line('SOFTWARE REQUIREMENTS SPECIFICATION', 11, after=38100)
cover_line('Manage Ticket Tracking · Sprint Board', 11)

pb = doc.add_paragraph()
pb.add_run().add_break(WD_BREAK.PAGE)

# ═════════════════════════════════════════════════════════════════════════════
# BODY
# ═════════════════════════════════════════════════════════════════════════════
h(1, 'Software Requirements Specification (SRS) — Manage Ticket Tracking')
normal('Feature: Manage Ticket Tracking (Sprint board) — Apex controller: '
       'force-app/main/default/classes/controller/manageTicketTracking/ManageTicketTrackingController.cls',
       bold=True)

# ── 1. Purpose & Scope ───────────────────────────────────────────────────────
h(2, '1. Purpose & Scope')
normal('The Manage Ticket Tracking page is the sprint-execution surface of the Jira Clone. '
       'For the single active sprint of a selected Project it lets a user:')
bullet('See the active sprint header (name, date range, goal) and a live story-point progress line.')
bullet('See one board column per project status, each holding the tickets currently in that status.')
bullet('Read every ticket as a card showing its name, ticket type, summary, story points, and assignee.')
bullet('Move a ticket from one status column to another by drag-and-drop, limited to the transitions allowed for that ticket’s type.')
bullet('Open a ticket in a centered ticket-view pop-up to edit its summary, status, and description.')
bullet('Manage a ticket’s linked work items, sub-tasks, history, and comments from the ticket-view pop-up.')
normal('The project is selected once and persisted in localStorage(’projectId’); if absent the page '
       'shows the Choose Project splash. Only the sprint whose status is in_progress is loaded; if there is '
       'none, a “No active sprint” banner is shown.')
figure('manage-ticket-tracking-default.png',
       'Figure 1 — The board: sprint header, story-point line, one column per status, and ticket cards. '
       'A ticket that has reached a Done (end) status shows a green left border.')

# ── 2. Actors ────────────────────────────────────────────────────────────────
h(2, '2. Actors')
actors_table([
    ('Actor', 'Description'),
    ('Project Member (User)',
     'The authenticated Salesforce user working the active sprint — moves tickets across status '
     'columns and edits ticket details (summary, status, description, links, sub-tasks, comments).'),
    ('System',
     'Apex services that enforce workflow transitions, run validation-field rules, log history, and keep '
     'the sprint’s story-point totals in sync.'),
])
doc.add_paragraph()

# ── 3. Functional Requirements ───────────────────────────────────────────────
h(2, '3. Functional Requirements')

h(3, '3.1 Project Selection and Data Loading')
bullet('FR-1 The system shall allow a user to select a project before accessing the ticket-tracking board.')
bullet('FR-2 The system shall load the board data for the selected project: statuses, workflow transitions, '
       'members, ticket types, the active sprint, and that sprint’s active tickets.')
bullet('FR-3 The system shall load only the sprint whose lifecycle state is in_progress; if none exists, it '
       'shall show a “No active sprint” banner instead of the board.')

h(3, '3.2 Sprint Header and Progress')
bullet('FR-4 The system shall display the active sprint’s name, start → end date range (end derived '
       'from start date + duration), and goal.')
bullet('FR-5 The system shall display a story-point progress line showing ended story points / total story '
       'points and the completion percentage.')

h(3, '3.3 Board Columns and Ticket Cards')
bullet('FR-6 The system shall render one column per project status; the column header shall show the status '
       'name and the count of tickets it holds.')
bullet('FR-7 Each column shall hold 0..many ticket cards — the tickets whose current status equals that '
       'column’s status — ordered by score.')
bullet('FR-8 Each ticket card shall show the ticket name, ticket type, summary, story points, and assigned '
       'project-member name.')
bullet('FR-9 A ticket that has reached an end (Done) status shall be marked with a green left border.')

h(3, '3.4 Ticket Movement (Status Change)')
bullet('FR-10 The system shall allow users to move a ticket to another status column by drag-and-drop.')
bullet('FR-11 On drag start the system shall highlight only the columns that are valid targets for that '
       'ticket, computed from the ticket type’s workflow transitions out of its current status (so a Task '
       'and a Story expose different reachable columns).')
bullet('FR-12 Dropping on a non-target column or on the ticket’s own column shall be ignored.')
bullet('FR-13 On drop the system shall apply the status change only if a workflow transition exists for '
       'fromStatus → toStatus and all validation-field rules attached to that transition pass; otherwise '
       'the move is rejected with a message and the card stays in place.')
bullet('FR-14 When a move takes a ticket to an end status, the sprint’s ended story points and the '
       'progress line shall update.')

h(3, '3.5 Ticket View (Pop-up)')
bullet('FR-15 The system shall allow users to open a ticket in a centered modal ticket-view pop-up by '
       'clicking its card.')
bullet('FR-16 The system shall allow users to close the ticket view and return to the board.')
bullet('FR-17 The system shall allow users to edit the ticket summary from the ticket view.')
bullet('FR-18 The system shall allow users to change ticket status from the ticket view, subject to the same '
       'workflow + validation-field rules as a board move.')
bullet('FR-19 The system shall allow users to edit the ticket description using a rich-text editor.')

h(3, '3.6 Linked Work Items')
bullet('FR-20 The system shall allow users to expand a ticket’s linked work items in the ticket view.')
bullet('FR-21 The system shall allow users to add a link by choosing a link type and searching for a target '
       'ticket (minimum 2 characters).')

h(3, '3.7 Sub-tasks')
bullet('FR-22 The system shall allow users to expand a ticket’s sub-tasks in the ticket view.')
bullet('FR-23 The system shall allow users to create a sub-task with summary (required) plus optional '
       'assignee, state, start date, story points, and description.')

h(3, '3.8 History and Comments')
bullet('FR-24 The system shall allow users to expand a ticket’s history (activity feed) in the ticket view.')
bullet('FR-25 The system shall allow users to expand and read a ticket’s comments.')
bullet('FR-26 The system shall allow users to add a comment to a ticket.')

figure('ticket-type-story-move.png',
       'Figure 2 — Dragging a Story ticket: the columns reachable from its current status, per the Story '
       'type’s workflow, are highlighted as drop targets (Drop here).')
figure('ticket-type-task-move.png',
       'Figure 3 — Dragging a Task ticket on the same board: a different set of columns is offered, '
       'because the Task type follows a different workflow.')

# ── 4. Validation Rules ──────────────────────────────────────────────────────
h(2, '4. Validation Rules')
h(3, '4.1 Client-side (LWC validators)')
bullet('VR-1 A status change requires a ticketId and a toStatusId (validateChangeTicketState) before any '
       'Apex call.')
bullet('VR-2 A drop is accepted only on a column flagged as a valid target; a drop on the ticket’s '
       'current column is a no-op.')
bullet('VR-3 A board move is allowed only when a matching workflow transition is found on the ticket type '
       '(findTransitionId); otherwise “This transition is not allowed by the workflow.” is shown.')
bullet('VR-4 Creating a sub-task requires a non-blank Summary.')
bullet('VR-5 Linking requires both a link type and a target ticket; ticket search ignores terms shorter than '
       '2 characters.')

h(3, '4.2 Server-side (controller input guards)')
bullet('VR-6 Every @AuraEnabled method blank-checks its required parameters and returns '
       'APIResponse(false, ’<field> is required’) before doing any work (e.g. projectId, ticketId, '
       'fromStatusId, toStatusId, summary, message, fromTicketId, toTicketId, linkType).')
bullet('VR-7 Referenced records must exist via DomainCorrectness.require* (project, ticket, ticket type, '
       'workflow, from/to status); optional references (assignee, sub-task state) use requireOptional*.')
bullet('VR-8 loadTicketBySearchTerm escapes all SOSL/SOQL special characters in the search term before '
       'building the query.')

# ── 5. Business Rules ────────────────────────────────────────────────────────
h(2, '5. Business Rules')
h(3, '5.1 Active sprint selection')
bullet('BR-1 The board loads exactly one sprint — the project’s sprint whose RecordStatus__c = '
       '’in_progress’ (at most one per project); if none, the board shows no sprint.')
bullet('BR-2 Only active tickets (RecordStatus__c = ’active’) of that sprint are loaded, ordered by '
       'Score__c ASC.')

h(3, '5.2 Columns & cards')
bullet('BR-3 Columns are derived from the project’s statuses (in their defined order); a ticket appears '
       'in the column whose status equals the ticket’s CurrentState__c.')
bullet('BR-4 A column can hold 0..many cards; the header count reflects the number of cards in that column.')
bullet('BR-5 A ticket whose current status has isEnd__c = true is an “ended” ticket and is rendered '
       'with a green left border.')

h(3, '5.3 Transitions & validation')
bullet('BR-6 A status change is allowed only if a WorkflowTransition__c exists for fromStatus → toStatus '
       'on the workflow of the ticket’s type (requireTransitionAllowed).')
bullet('BR-7 Every ValidateField__c rule attached to that transition must pass (requireValidateFieldsPass) '
       'before the status is written; a failing rule rejects the change with the rule’s error message.')
bullet('BR-8 The set of reachable columns is therefore ticket-type specific — two types pointing at '
       'different workflows expose different valid targets from the same status.')
bullet('BR-9 An end status is terminal: no transition may originate from it.')

h(3, '5.4 Story-point accounting')
bullet('BR-10 The sprint’s story-point line shows TotalEndedStoryPoint__c / TotalStoryPoint__c and a '
       'percentage = round(ended / total × 100).')
bullet('BR-11 When a ticket in the sprint reaches an end status, its story points are added to the sprint’s '
       'ended total and the progress line refreshes.')

h(3, '5.5 Data hygiene')
bullet('BR-12 All ticket, sub-task, link, history, and comment reads exclude soft-deleted rows '
       '(RecordStatus__c = ’active’ / != ’deleted’), per the project’s '
       'soql-exclude-deleted rule.')

# ── 6. Use Case Descriptions ─────────────────────────────────────────────────
h(2, '6. Use Case Descriptions')

h(3, 'UC-1 — Load the board')
bullet('Actor: Project Member')
bullet('Pre-conditions: User is authenticated; a projectId exists in localStorage.')
bullet('Main flow:')
normal('1. Component mounts and reads projectId. 2. Calls loadManageTicketTrackingPage(projectId). '
       '3. Renders the sprint header, the story-point line, and one column per status filled with the active '
       'sprint’s tickets.')
bullet('Alternative flows:')
bullet('A1 (no project): No projectId → Choose-Project splash; on projectchosen the project is stored and '
       'the main flow resumes from step 2.')
bullet('A2 (no active sprint): The sprint is null → “No active sprint. Start a sprint from the '
       'Backlog page.” banner; columns are not shown.')
bullet('A3 (load error): Apex returns success=false or throws → error toast, content hidden.')

h(3, 'UC-2 — Move a ticket between columns')
bullet('Actor: Project Member')
bullet('Pre-conditions: The board is loaded with an active sprint and at least two statuses.')
bullet('Main flow:')
normal('1. User starts dragging a ticket card. 2. The system highlights the columns reachable from the '
       'ticket’s current status for that ticket type. 3. User drops on a highlighted column. '
       '4. changeTicketState(ticketId, fromStatusId, toStatusId) validates the transition and its '
       'validation-field rules, writes the new status, and — if the new status is an end status — '
       'updates the sprint’s ended story points. 5. The card moves to the target column.')
bullet('Alternative flows:')
bullet('A1 (non-target column): Drop is ignored; the card stays.')
bullet('A2 (same column): No-op.')
bullet('A3 (transition not allowed / validation-field fails): Error toast; the card stays in place.')

figure('ticket-view-pop-up.png',
       'Figure 4 — Clicking a ticket card opens the ticket view as a pop-up centered on the screen, '
       'showing the TKT-id, title, status, description, and the Linked work items / Subtasks / History / '
       'Comments sections.')

h(3, 'UC-3 — Open the ticket view')
bullet('Actor: Project Member')
bullet('Pre-conditions: A ticket card is visible on the board.')
bullet('Main flow:')
normal('1. User clicks a ticket card. 2. ticketviewopen records the active ticket Id. 3. A centered modal '
       'renders the ticket: its TKT-id, summary, current status, description, and the Linked work items / '
       'Sub-tasks / History / Comments sections.')
bullet('Alternative flows:')
bullet('A1 (close): Clicking the backdrop or the × clears the active ticket and unmounts the panel.')

h(3, 'UC-4 — Edit summary, status, or description from the ticket view')
figure('ticket-view-click-on-description.png',
       'Figure 5 — Clicking the description opens a rich-text editor with Save / Cancel.')
bullet('Actor: Project Member')
bullet('Pre-conditions: The ticket view is open.')
bullet('Main flow (any one of):')
bullet('Edit summary → ticketsummaryupdate → updateTicketSummary.')
bullet('Change status (combo) → ticketstatuschange → changeTicketState (BR-6 / BR-7).')
bullet('Edit description → click the description → rich-text editor → Save → '
       'ticketdescriptionupdate → updateTicketDescription.')
bullet('Alternative flows:')
bullet('A1 (server error on any action): Error toast; the panel keeps its last good value.')

h(3, 'UC-5 — Link work items')
figure('ticket-view-expand-ticket-view.png',
       'Figure 6 — Expanding Linked work items lists the ticket’s links (BLOCKS); empty shows '
       '“No linked items.”')
figure('ticket-ciew-click-plus.png',
       'Figure 7 — Clicking + reveals the add-link row: choose a Link type and search a target ticket '
       '(min 2 chars), then Link.')
bullet('Actor: Project Member')
bullet('Pre-conditions: The ticket view is open.')
bullet('Main flow:')
normal('1. User expands Linked work items (ticketlinkedtoexpand → loadTicketLinkedTo lists existing '
       'links). 2. User clicks +, chooses a link type, and searches a target ticket (ticketsearch → '
       'loadTicketBySearchTerm). 3. User clicks Link → linkToTicket adds the link to the list.')
bullet('Alternative flows:')
bullet('A1 (incomplete): Link stays disabled until both a type and a target ticket are chosen; searches '
       'under 2 characters are ignored.')

h(3, 'UC-6 — Create and expand sub-tasks')
figure('ticket-view-expand-subtask.png',
       'Figure 8 — Expanding Subtasks lists the ticket’s sub-tasks; empty shows “No subtasks.”')
figure('ticket-view-click-plus-subtask.png',
       'Figure 9 — Clicking + opens the create-subtask form: Summary (required), Assignee, Current State, '
       'Start Date, Story Point, and a rich-text Description.')
bullet('Actor: Project Member')
bullet('Pre-conditions: The ticket view is open.')
bullet('Main flow:')
normal('1. User expands Sub-tasks (subtasksexpand → loadSubtasks). 2. User clicks +, fills the form '
       '(Summary required; optional assignee, state, start date, story points, description). 3. Create '
       '→ createSubtask adds the sub-task to the list.')
bullet('Alternative flows:')
bullet('A1 (summary blank): Required-field validation; no call.')
bullet('A2 (cancel): The form closes.')

h(3, 'UC-7 — View ticket history')
figure('ticket-view-expand-history.png',
       'Figure 10 — Expanding History shows the ticket’s activity feed (e.g. “Ticket created”) '
       'with author and timestamp.')
bullet('Actor: Project Member')
bullet('Pre-conditions: The ticket view is open.')
bullet('Main flow:')
normal('1. User expands History (tickethistoryexpand → loadTicketHistory). 2. The system lists the '
       'ticket’s activity feed, each entry with its author and timestamp.')

h(3, 'UC-8 — Read and add comments')
figure('ticket-view-expand-comment.png',
       'Figure 11 — Expanding Comments shows the add-comment box and the existing comments, each with '
       'author and timestamp.')
bullet('Actor: Project Member')
bullet('Pre-conditions: The ticket view is open.')
bullet('Main flow:')
normal('1. User expands Comments (ticketcommentsexpand → loadTicketComments). 2. User types a comment '
       'and clicks Comment → createTicketComment prepends it to the list.')
bullet('Alternative flows:')
bullet('A1 (re-expand): Re-expanding the same ticket refreshes the comments via refreshApex.')

# ── 7. Non-Functional Notes ──────────────────────────────────────────────────
h(2, '7. Non-Functional Notes')
bullet('Read efficiency: the page aggregate, linked-items, search, sub-tasks, history, and comments reads are '
       'cacheable=true; linked items, sub-tasks, history, and comments load lazily on expand.')
bullet('Derived UI state: the LWC keeps one principal state (the sprint with its tickets) and derives every '
       'displayed value — columns, valid drop targets, the active ticket-view model, the story-point '
       'percentage — via getters and formatter utilities, never parallel tracked fields.')
bullet('Type-aware transitions client-side: valid drop targets are computed in the browser from each ticket '
       'type’s nested workflow transitions, so no round-trip is needed to highlight reachable columns.')
bullet('Soft delete everywhere: tickets, sub-tasks, links, and comments are soft-deleted; queries filter '
       'them out.')

# ── save ─────────────────────────────────────────────────────────────────────
doc.save(OUT)
if os.path.exists(TMP_LOGO):
    os.remove(TMP_LOGO)
print('WROTE', OUT)
