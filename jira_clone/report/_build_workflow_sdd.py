# -*- coding: utf-8 -*-
"""
Build report/final/SDD_ManageWorkflow.docx.

Reuses report/final/SDD_ManageBacklog.docx as the *style + structure template*:
same cover, Aptos Display headings, API-method tables (Kind C/I legend),
Consolas ASCII architecture / sequence diagrams, the ###-wrapped image marker
convention, and the Traceability table. Content is grounded in the actual
manageWorkflow LWC + ManageWorkflowPageController code.

This feature has no screenshots, so — per the request — each image marker pair
encloses a one-line italic comment describing the image that should be placed
there (instead of an embedded picture). Because no images are embedded, the
orphan-media prune at the end strips the template screenshots so the output only
carries the cover logo.
"""
import os, zipfile
from docx import Document
from docx.shared import Pt, Emu, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT      = os.path.dirname(os.path.abspath(__file__))
TEMPLATE  = os.path.join(ROOT, 'final', 'SDD_ManageBacklog.docx')
OUT       = os.path.join(ROOT, 'final', 'SDD_ManageWorkflow.docx')
TMP_LOGO  = os.path.join(ROOT, '_cover_logo.png')

# extract the cover logo (word/media/image1.png) from the template
with zipfile.ZipFile(TEMPLATE) as z:
    with open(TMP_LOGO, 'wb') as f:
        f.write(z.read('word/media/image1.png'))

doc = Document(TEMPLATE)

# wipe the template body, keep the trailing sectPr (page setup)
body = doc.element.body
for child in list(body):
    if child.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(child)

CONTENT_W     = Inches(6.3)
HEADING_COLOR = RGBColor(0x0F, 0x47, 0x61)
CAPTION_GREY  = RGBColor(0x6B, 0x77, 0x8C)
CENTER        = WD_ALIGN_PARAGRAPH.CENTER
HASH          = '#' * 27

# ── helpers ───────────────────────────────────────────────────────────────────
def set_consolas(run, size):
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(a), 'Consolas')

def cover_line(text, size, italic=True, black=False, before=None, after=None):
    p = doc.add_paragraph(); p.alignment = CENTER
    if before is not None: p.paragraph_format.space_before = Emu(before)
    if after  is not None: p.paragraph_format.space_after  = Emu(after)
    r = p.add_run(text)
    r.italic = italic
    r.font.name = 'Aptos Display'
    r.font.size = Pt(size)
    if black:
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def h(level, text):
    return doc.add_heading(text, level=level)

def normal(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold
    return p

def label(text):
    """Bold inline label (used for the 3.x sequence-flow headers)."""
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True
    return p

def bullet(text):
    return doc.add_paragraph('•  ' + text)

def mono(text, size=9):
    """Consolas block paragraph; '\\n' becomes a real line break (w:br)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        run = p.add_run(ln); set_consolas(run, size)
        if i != len(lines) - 1:
            br = p.add_run(); br.add_break(WD_BREAK.LINE)
    return p

def marker(filename):
    """One image marker block: ### / filename / ### (Consolas 8pt)."""
    for line in (HASH, filename, HASH):
        p = doc.add_paragraph()
        set_consolas(p.add_run(line), 8)

def figure_comment(filename, desc):
    """Marker convention with NO embedded image: the filename is wrapped before
    and after a one-line italic comment describing the image to be placed."""
    marker(filename)
    p = doc.add_paragraph(); p.alignment = CENTER
    r = p.add_run('Image to place — ' + desc)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = CAPTION_GREY
    marker(filename)

def _table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom', 'insideH'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0');    el.set(qn('w:color'), 'D0D7DE')
        borders.append(el)
    tblPr.append(borders)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), '9360'); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)

def api_table(headers, rows, mono_cols=(0,), center_cols=()):
    """API-method table. col 0 of mono_cols renders Consolas 10pt, the rest 9pt."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    for j, head in enumerate(headers):
        cell = table.rows[0].cells[j]; cell.text = head
        for r in cell.paragraphs[0].runs: r.bold = True
        if j in center_cols: cell.paragraphs[0].alignment = CENTER
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]; cell.text = val
            para = cell.paragraphs[0]
            if j in mono_cols:
                for r in para.runs:
                    set_consolas(r, 10 if j == 0 else 9)
            if j in center_cols:
                para.alignment = CENTER
    _table_borders(table)
    return table

def trace_table(rows):
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (a, b) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text, c1.text = a, b
        c0.width = Emu(2971800); c1.width = Emu(2971800)
        if i == 0:
            for c in (c0, c1):
                for r in c.paragraphs[0].runs: r.bold = True
        else:
            for r in c1.paragraphs[0].runs: set_consolas(r, 9)
    _table_borders(table)
    return table

# ═════════════════════════════════════════════════════════════════════════════
# COVER
# ═════════════════════════════════════════════════════════════════════════════
img_p = doc.add_paragraph(); img_p.alignment = CENTER
img_p.paragraph_format.space_before = Emu(1524000)
img_p.add_run().add_picture(TMP_LOGO, width=Inches(3.12))

cover_line('Technical Document', 24, black=True, before=381000, after=76200)
cover_line('Manage Workflow — SDD', 18, black=True, after=254000)
cover_line('SOFTWARE DESIGN DOCUMENT', 11, after=38100)
cover_line('Manage Workflow · Workflow Designer', 11)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ═════════════════════════════════════════════════════════════════════════════
# BODY
# ═════════════════════════════════════════════════════════════════════════════
h(1, 'Software Design Document (SDD) — Manage Workflow')

# feature line (bold labels, plain paths) — mirrors SDD_ManageBacklog
fp = doc.add_paragraph()
def seg(t, bold=False):
    r = fp.add_run(t); r.bold = bold
seg('Feature:', True); seg(' Manage Workflow (Workflow designer) ')
seg('Apex controller:', True)
seg(' force-app/main/default/classes/controller/manageWorkflow/ManageWorkflowPageController.cls ')
seg('LWC:', True); seg(' force-app/main/default/lwc/manageWorkflow/ ')
seg('Date:', True); seg(' 2026-06-22')

normal('Image references below follow the marker convention used across these design documents: the image '
       'file name is wrapped between ########################### lines, immediately before and after where '
       'the screenshot belongs. This document ships without the screenshots — so each marker pair encloses a '
       'one-line comment describing the image that should be placed there instead of an embedded image.')

# ── 1. Architecture Overview ──────────────────────────────────────────────────
h(2, '1. Architecture Overview')

ARCH = (
"┌─ manageWorkflow (parent LWC) ───────────────────────────────────────────────┐\n"
"│ Principal state: workflowData = {                                            │\n"
"│   id, projectStatus:[ { id, name } ],                                        │\n"
"│   workflow:{ id, name, transitions:[                                         │\n"
"│     { id, name, fromStatus, toStatus, recordStatus, validateFields[] } ] } } │\n"
"│ UI-only state (outside principal): _projectId · _workflowId ·               │\n"
"│   clickedStatusIds · selectedFromStatus/selectedToStatus ·                  │\n"
"│   selectedTransitionId · _isValidationDetailExpanded · _wiredTransitionId · │\n"
"│   config (responsive) · modal flags (status/transition/workflow/validation) │\n"
"│ Derived getters: sortedStatuses · statusPositions · statusesWithSVGData ·   │\n"
"│   activeTransitionLines · pendingTransitionLines · svgViewBox ·             │\n"
"│   activeTransition · validationFields · showTransitionDetail               │\n"
"│                                                                            │\n"
"│ workflowUtils.js (pure): getResponsiveConfig · gatherSortedStatuses ·      │\n"
"│   calculatePositions · calculateTransitionLines · createArrowPath ·        │\n"
"│   getRectIntersection · getSvgViewBox · getStatusesWithSVGData · toggleClick│\n"
"│ workflowValidator.js (pure): validateStatusName · validateTransition ·     │\n"
"│   validateTransitionName · validateValidationType · validateTicketField ·  │\n"
"│   validateTransitionId                                                     │\n"
"│                                                                            │\n"
"│ Children                                                                    │\n"
"│   <c-choose-project> (step 1, project picker)                              │\n"
"│   workflow list (step 2): cards + Create-Workflow modal                    │\n"
"│   SVG canvas (step 3): status rects · transition arrows (active/pending)  │\n"
"│   modals: create status · create transition · transition-detail panel ·   │\n"
"│     add-validation-rule section                                            │\n"
"└──────────────────────────────────────────────────────────────────────────────┘\n"
"        │  imperative writes/reads  +  @wire(loadValidateFields, cacheable=true)\n"
"        ▼\n"
"  ManageWorkflowPageController\n"
"     → WorkflowService · StatusService · ValidateFieldService · DomainCorrectness\n"
"     → Workflow__c · Status__c · WorkflowTransition__c · ValidateField__c · Project__c"
)
mono(ARCH, size=9)

normal('Response envelope. Every controller method returns APIResponse { success, message, data }. '
       'The LWC treats success === false (or a thrown Apex error) as failure and toasts it via '
       '_toast(...,"error"); successful writes patch the principal state from the returned data (the '
       'persisted record), never from an optimistic draft.')
normal('Data shape used by the editor (from getWorkflow → WorkflowConfigDTO): '
       '{ id, projectStatus:[{id,name}], workflow:{ id, name, transitions:[{id,name,fromStatus,toStatus,'
       'recordStatus}] } }. A transition’s validateFields[] are not part of this payload — they are loaded '
       'lazily per transition by loadValidateFields and folded onto the owning transition inside the same '
       'principal state, so the canvas needs no further round-trip to draw the graph.')

# ── 2. API Calls ──────────────────────────────────────────────────────────────
h(2, '2. API Calls (Apex @AuraEnabled methods)')
lg = doc.add_paragraph()
lg.add_run('Legend: ')
r = lg.add_run('C'); r.bold = True
lg.add_run(' = cacheable=true (wire-friendly), ')
r = lg.add_run('I'); r.bold = True
lg.add_run(' = imperative (cacheable=false or DML). All methods return APIResponse{success,message,data}.')

h(3, '2.1 Workflow list & editor bootstrap')
api_table(
    ['Method', 'Kind', 'Params', 'Returns (`data`)', 'LWC usage'],
    [
        ['loadWorkflowsByProject', 'C', 'projectId',
         'Workflow__c[] (Id, Name, RecordStatus__c, dates)',
         '_loadWorkflowsForProject() — called imperatively to build the list'],
        ['getWorkflow', 'I', 'workflowId',
         'WorkflowConfigDTO {projectStatus, workflow{transitions}}',
         '_loadWorkflow() on entering the editor (cacheable=false)'],
        ['createWorkflow', 'I', 'name, projectId', 'Workflow__c',
         'Create-Workflow modal → then opens editor on the new id'],
    ],
    mono_cols=(0, 2, 3), center_cols=(1,))

h(3, '2.2 Status (node) creation')
api_table(
    ['Method', 'Kind', 'Params', 'Effect'],
    [
        ['createStatus', 'I', 'name, projectId',
         'creates a project Status__c; returns the record → _addStatus({id,name})'],
    ],
    mono_cols=(0, 2, 3), center_cols=(1,))

h(3, '2.3 Transition create / lifecycle')
api_table(
    ['Method', 'Kind', 'Params', 'Effect'],
    [
        ['addWorkflowTransition', 'I', 'workflowId, name, fromStatusId, toStatusId',
         'creates a WorkflowTransition__c in pending status; returns it → _addPendingTransition(...)'],
        ['activateWorkflowTransition', 'I', 'workflowTransitionId',
         'pending → active; LWC → _setTransitionRecordStatus(id,"active")'],
        ['deleteWorkflowTransition', 'I', 'workflowTransitionId',
         'soft delete (record status → deleted); LWC → _removeTransition(id)'],
        ['updateFullWorkflowTransition', 'I', 'workflowId',
         'activates all pending transitions of the workflow; returns the activated count'],
    ],
    mono_cols=(0, 2, 3), center_cols=(1,))

h(3, '2.4 Validation rules')
api_table(
    ['Method', 'Kind', 'Params', 'Effect'],
    [
        ['loadValidateFields', 'C', 'transitionId',
         'ValidateField__c[] for the transition (wire, on expand)'],
        ['addValidateField', 'I', 'workflowTransitionId, fieldName, type',
         'creates a ValidateField__c; returns it → _addTransitionValidateField(id, ...)'],
    ],
    mono_cols=(0, 2, 3), center_cols=(1,))

h(3, '2.5 Wires actually bound in manageWorkflow.js')
bullet('@wire(loadValidateFields, { transitionId: \'$_wiredTransitionId\' }) → maps each ValidateField__c and '
       'folds the list onto the open transition via _setTransitionValidateFields. The gate _wiredTransitionId '
       'stays undefined (wire dormant) until the user expands the detail, so selecting a transition does not by '
       'itself fetch its rules; collapsing clears the gate so the next expand re-fires.')
normal('All other interactions — the workflow list load, getWorkflow, and every write (createWorkflow, '
       'createStatus, addWorkflowTransition, activateWorkflowTransition, deleteWorkflowTransition, '
       'updateFullWorkflowTransition, addValidateField) — are imperative Apex calls inside handlers '
       '(then/catch/finally), toasting via _toast.')
normal('The controller also exposes endpoints this LWC does not bind (getProjectStatuses, loadWorkflows, '
       'loadStatuses, loadTransitions, getWorkflowTransitionById, getTicketTypeById, updateWorkflow). They are '
       'available for reuse but out of scope for the Manage Workflow page.')

# ── 3. Sequence Flows ─────────────────────────────────────────────────────────
h(2, '3. Sequence Flows (key interactions)')

label('3.1 Page load & three-step entry')
mono(
"connectedCallback → read localStorage.projectId\n"
" ├─ none     → render <c-choose-project>; on projectchosen → _projectId → _loadWorkflowsForProject()\n"
" └─ present  → _loadWorkflowsForProject() → loadWorkflowsByProject(projectId)\n"
"       success → _workflows = mapped rows → workflow list (cards + Create New Workflow)\n"
"       empty   → \"No workflows for this project yet\" empty-state\n"
"       fail    → _workflowsErrorMessage + error toast\n"
" Edit a workflow → _enterWorkflowEditor(id):\n"
"   _workflowId = id ; localStorage.workflowId = id → _loadWorkflow() → getWorkflow(id)\n"
"     success → workflowData set → derived getters render the SVG graph\n"
"     fail    → errorMessage + error toast")

label('3.2 Create a status')
mono(
"Create Status (toolbar) → openCreateStatusModal\n"
"  submit → validateStatusName(name)  (required, ≤80)\n"
"         → createStatus(name, projectId)\n"
"             success → _addStatus({id,name}) (immutable write to workflowData.projectStatus)\n"
"                     → clickedStatusIds = clearClicks() ; close modal\n"
"             fail    → createStatusErrorMessage + error toast")

label('3.3 Create a transition (two-click status selection)')
mono(
"status click (handleStatusClick) →\n"
"   no selectedFromStatus            → selectedFromStatus = status\n"
"   click same status again          → selectedFromStatus = null (deselect)\n"
"   click a different status         → selectedToStatus = status → openCreateTransitionModal()\n"
"   (every click toggles clickedStatusIds for the highlight)\n"
"submit (handleCreateTransitionSubmit) →\n"
"   validateTransition({from,to})  (both present, distinct)\n"
"   validateTransitionName(name)\n"
"   addWorkflowTransition(workflowId, name, fromStatusId, toStatusId)\n"
"     success → _addPendingTransition({id,name,fromStatus,toStatus, recordStatus:'pending'})\n"
"             → pendingTransitionLines getter draws a dashed arrow ; clearClicks ; close\n"
"     fail    → createTransitionErrorMessage + error toast")

label('3.4 Transition detail — activate / delete')
mono(
"arrow click (handleTransitionClick) → selectedTransitionId = lineId\n"
"   → showTransitionDetail (derived: id != null) ; activeTransition = _transitions.find(id)\n"
"   → _collapseValidationDetail() (drop any open rules from the previous transition)\n"
"Activate (pending only) → activateWorkflowTransition(id)\n"
"   success → _setTransitionRecordStatus(id,'active') → arrow re-rendered solid ; close panel\n"
"Delete → confirm() → deleteWorkflowTransition(id)\n"
"   success → _removeTransition(id) → arrow disappears ; close panel\n"
"Close → handleCloseTransitionDetail → selectedTransitionId = null (+ reset panel state)")

label('3.5 Validation rules — add + expand (wire)')
mono(
"Add Validation Rule → openValidationRuleModal (inline section in the panel)\n"
"  submit → validateTicketField(fieldName) ; validateValidationType(type)\n"
"         → addValidateField(transitionId, fieldName, type)\n"
"             success → _addTransitionValidateField(id, mapped) (folded onto the transition)\n"
"Show Validation Details (handleToggleValidationDetail) →\n"
"  expanding → validateTransitionId(selectedTransitionId)\n"
"            → _isValidationDetailExpanded = true ; _wiredTransitionId = selectedTransitionId\n"
"            → @wire(loadValidateFields) fires → _setTransitionValidateFields(id, mapped)\n"
"  collapsing → _isValidationDetailExpanded = false ; _wiredTransitionId = undefined")

label('3.6 Update workflow (activate all pending)')
mono(
"Update Workflow (toolbar) → handleUpdateWorkflow\n"
"  guard _workflowId → updateFullWorkflowTransition(workflowId)\n"
"    success → toast count ; _activatePendingTransitions() (local: pending → active)\n"
"            → handleBackToWorkflowList() (clear workflowId, teardown ResizeObserver, reload list)\n"
"    fail    → errorMessage + error toast")

label('3.7 Responsive re-layout (ResizeObserver)')
mono(
"renderedCallback → observe('.workflow-visualizer-container')\n"
"  on resize → if |width − _lastMeasuredWidth| > 20:\n"
"      _lastMeasuredWidth = width ; config = getResponsiveConfig(width)  (immutable reassign)\n"
"      → statusPositions / svgViewBox / *TransitionLines getters re-run with the new geometry\n"
"disconnectedCallback / Back → _teardownResizeObserver() → config reset to VISUALIZATION_CONFIG")

# ── 4. UI / UX Specification ──────────────────────────────────────────────────
h(2, '4. UI / UX Specification')

h(3, '4.1 Step 1 — project picker')
normal('If there is no projectId the page renders <c-choose-project>; on projectchosen the project is stored '
       'and the workflow list loads.')
figure_comment('manage-workflow-choose-project.png',
               'The Choose Project splash shown when no project is selected (the <c-choose-project> child).')

h(3, '4.2 Step 2 — workflow list')
normal('With a project but no workflow selected the page shows the “Workflows” heading, a Create New Workflow '
       'button, and one card per workflow (name + record status) each with an Edit button; an empty project '
       'shows “No workflows for this project yet — create one to get started.” The Create New Workflow modal '
       'takes a name and, on success, opens the editor on the new workflow.')
figure_comment('manage-workflow-list.png',
               'The workflow list: “Workflows” heading, “Create New Workflow” button, and workflow cards '
               '(name + record status + Edit).')
figure_comment('manage-workflow-create-modal.png',
               'The Create New Workflow modal: a Workflow Name input with Cancel / Create.')

h(3, '4.3 Step 3 — the visual editor (toolbar, legend, canvas)')
normal('The editor renders a toolbar (Back · Create Status · Update Workflow) and an info span that reads '
       '“Click two statuses to create a transition”, or “From: X → To: Y” once a source is selected. A legend '
       'distinguishes Active vs Pending transitions. The canvas is a single SVG whose viewBox is derived from '
       'the status count and the responsive config.')
figure_comment('manage-workflow-editor.png',
               'The editor: toolbar, Active / Pending legend, status nodes in a grid, and curved transition '
               'arrows (solid = active, dashed = pending) with name labels.')

h(3, '4.4 Status nodes & selection')
normal('Each status is a rounded rectangle labelled with the status name; clicking a node selects it '
       '(highlighted via clickedStatusIds). The first click sets the From status; clicking it again deselects; '
       'clicking a different node sets the To status and opens the Create Transition dialog. Create Status adds '
       'a new node.')
figure_comment('manage-workflow-status-selected.png',
               'A status node selected as “From” (highlighted), with the toolbar info span showing '
               '“From: <status>”.')

h(3, '4.5 Transition arrows (active vs pending)')
normal('Transitions are drawn as quadratic-Bézier curved arrows between node edges (getRectIntersection finds '
       'where the curve meets each rectangle; createArrowPath curves it to avoid overlapping nodes). Active '
       'transitions are solid; pending transitions are dashed. Each arrow carries the transition name at its '
       'mid-point and an explicit arrow-head polygon at the target end.')
figure_comment('manage-workflow-create-transition.png',
               'The Create Transition dialog after selecting two nodes: “From → To” badges and a Transition '
               'Name field with Cancel / Create Transition.')

h(3, '4.6 Transition detail panel')
normal('Clicking an arrow opens the detail panel: Name, Id, From / To status badges, Record Status badge, and '
       'Created Date. While the transition is pending an Activate action is shown; Delete (with a confirm '
       'prompt) soft-deletes it; Close dismisses the panel. The panel also hosts Add Validation Rule and Show '
       '/ Hide Validation Details.')
figure_comment('manage-workflow-transition-detail.png',
               'The transition detail panel: Name / Id / From / To / Record Status / Created Date, with '
               'Activate (pending) · Delete · Close and the Add Validation Rule / Show Validation Details '
               'controls.')

h(3, '4.7 Validation rules')
normal('Add Validation Rule opens an inline section with a Field API Name combobox (mirroring the '
       'ValidateField__c field picklist) and a Validation Type combobox; Create persists the rule and folds it '
       'onto the transition. Show Validation Details expands the rule list (loaded on expand via '
       'loadValidateFields); an empty transition shows “No validation details for this transition.”')
figure_comment('manage-workflow-add-validation-rule.png',
               'The Add Validation Rule inline section: Field API Name and Validation Type comboboxes with '
               'Cancel / Create.')
figure_comment('manage-workflow-validation-details.png',
               'A transition expanded to show its validation rules (field · type · error message rows).')

# ── 5. State, Geometry & Derived View Design ──────────────────────────────────
h(2, '5. State, Geometry & Derived View Design')
bullet('Principal state. One de-normalized workflowData = { id, projectStatus[], workflow{ id, name, '
       'transitions[] } }. Everything the editor shows is a getter derived on read — sortedStatuses, '
       'statusPositions, statusesWithSVGData, activeTransitionLines, pendingTransitionLines, svgViewBox, '
       'activeTransition, validationFields — never a parallel tracked field.')
bullet('Accessors & mutators only. The workflow is read through _statuses / _transitions and written only '
       'through named immutable mutators — _addStatus, _addPendingTransition, _setTransitionRecordStatus, '
       '_activatePendingTransitions, _removeTransition, _setTransitionValidateFields, '
       '_addTransitionValidateField, all funnelling through _writeTransitions. There is no inline '
       'workflowData.workflow.transitions mutation anywhere else.')
bullet('Folded detail. A transition’s validateFields live on that transition inside the principal state (set '
       'when its detail loads, appended when one is created) — not in a throwaway panel buffer, so '
       're-selecting a transition keeps its already-loaded rules.')
bullet('Derived geometry. gatherSortedStatuses builds the node list (project statuses plus any status only '
       'referenced by a transition); calculatePositions lays them on a grid; calculateTransitionLines computes '
       'each curved arrow’s path, arrow-head polygon, and label mid-point; getSvgViewBox sizes the canvas. All '
       'are pure functions of workflowData + config.')
bullet('Two derived line sets. activeTransitionLines / pendingTransitionLines filter _transitions by record '
       'status and feed the same geometry function, so a transition’s appearance (solid vs dashed) follows its '
       'state with no duplicated drawing logic.')
bullet('Responsive config. config is reassigned immutably by the ResizeObserver (getResponsiveConfig returns a '
       'new object); because the geometry getters read config, one assignment re-flows the whole graph '
       '(4 / 3 / 2 columns at > 768 / ≤ 768 / ≤ 480 px). Back / unmount tears the observer down and resets '
       'config.')
bullet('Separate wire gate. selectedTransitionId is the principal UI selection; _wiredTransitionId is a '
       'distinct gate set only when the user expands the detail, so the loadValidateFields wire fires on '
       'expand — not on every selection.')

# ── 6. Error Handling & Loading ───────────────────────────────────────────────
h(2, '6. Error Handling & Loading')
bullet('isLoading gates the editor (a c-ao-spinner shows while _loadWorkflow runs); _workflowsLoading gates '
       'the list; workflowListShouldShowContent hides the list while loading or on error.')
bullet('Every imperative call checks res.success — on false it surfaces res.message via a toast and aborts the '
       'patch; thrown Apex errors are caught and surfaced via err.body?.message || err.message || "<fallback>".')
bullet('Page / list load failures set errorMessage / _workflowsErrorMessage and toast, so the canvas is not '
       'shown with partial data.')
bullet('Server guards mirror the client: each @AuraEnabled method blank-checks required params and returns '
       'APIResponse(false, "<field> is required"), then DomainCorrectness.require* verifies referenced records '
       '(project, workflow, status, transition) before any DML.')

# ── 7. Traceability (SRS → API) ───────────────────────────────────────────────
h(2, '7. Traceability (SRS → API)')
trace_table([
    ('SRS', 'API call(s) / mechanism'),
    ('FR-1 select project', '<c-choose-project> → localStorage projectId'),
    ('FR-2/FR-3 workflow list & empty state', 'loadWorkflowsByProject → _workflows / empty-state getter'),
    ('FR-4 create workflow', 'createWorkflow → _enterWorkflowEditor'),
    ('FR-5/FR-6 open editor & load config', 'getWorkflow (WorkflowConfigDTO)'),
    ('FR-7 back to list', 'handleBackToWorkflowList (+ ResizeObserver teardown)'),
    ('FR-8 render status nodes', 'gatherSortedStatuses → statusesWithSVGData (derived)'),
    ('FR-9 create status', 'createStatus → _addStatus'),
    ('FR-10 responsive layout', 'getResponsiveConfig via ResizeObserver'),
    ('FR-11/FR-12 select two statuses & dialog', 'handleStatusClick → openCreateTransitionModal'),
    ('FR-13 pending transition', 'addWorkflowTransition → _addPendingTransition'),
    ('FR-14/FR-16 arrows & labels', 'calculateTransitionLines (active vs pending sets)'),
    ('FR-15 legend', 'static legend (active / pending)'),
    ('FR-17 transition detail', 'handleTransitionClick → activeTransition (derived)'),
    ('FR-18 activate', 'activateWorkflowTransition → _setTransitionRecordStatus'),
    ('FR-19 delete', 'deleteWorkflowTransition (soft) → _removeTransition'),
    ('FR-20 close panel', 'handleCloseTransitionDetail'),
    ('FR-21 add validation rule', 'addValidateField → _addTransitionValidateField'),
    ('FR-22 view validation rules', 'loadValidateFields (wire, on expand) → _setTransitionValidateFields'),
    ('FR-23 update workflow', 'updateFullWorkflowTransition → _activatePendingTransitions'),
])

# ── save ──────────────────────────────────────────────────────────────────────
doc.save(OUT)
if os.path.exists(TMP_LOGO):
    os.remove(TMP_LOGO)


def prune_orphan_media(path):
    """Remove image parts (and their relationships) the body no longer references."""
    import re, shutil
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        referenced = set()
        rid_re = re.compile(r'r:(?:embed|link|id)="(rId\d+)"')
        for n in names:
            if n.startswith('word/') and n.endswith('.xml') and '/_rels/' not in n:
                referenced.update(rid_re.findall(z.read(n).decode('utf-8', 'ignore')))
        rels_name = 'word/_rels/document.xml.rels'
        rels = z.read(rels_name).decode('utf-8')
        rel_re = re.compile(r'<Relationship\b[^>]*?/>')
        drop_media = set()
        new_rels = rels
        for tag in rel_re.findall(rels):
            rid = re.search(r'Id="([^"]+)"', tag).group(1)
            typ = re.search(r'Type="([^"]+)"', tag).group(1)
            tgt = re.search(r'Target="([^"]+)"', tag).group(1)
            if typ.endswith('/image') and rid not in referenced:
                drop_media.add('word/' + tgt.replace('../', '').lstrip('/'))
                new_rels = new_rels.replace(tag, '')
        if not drop_media:
            return 0
        payload = {n: z.read(n) for n in names if n not in drop_media}
        payload[rels_name] = new_rels.encode('utf-8')
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
        for n, data in payload.items():
            zo.writestr(n, data)
    shutil.move(tmp, path)
    return len(drop_media)


removed = prune_orphan_media(OUT)
print('WROTE', OUT, '| pruned', removed, 'orphan image(s)')
