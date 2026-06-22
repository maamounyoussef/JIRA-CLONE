# -*- coding: utf-8 -*-
"""
Build report/final/SRS_ManageWorkflow.docx.

Reuses report/final/SRS_ManageBacklog.docx as a *style template* (same cover
layout, Aptos Display headings, table look) so the new SRS matches the format and
composition of the SRS series exactly. Content is grounded in the actual
manageWorkflow LWC + ManageWorkflowPageController code.

This feature has no screenshots, so — per the request — every figure is rendered
as a one-line italic caption describing the image that should be placed there
(no embedded picture).
"""
import os, zipfile
from docx import Document
from docx.shared import Pt, Emu, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT      = os.path.dirname(os.path.abspath(__file__))
TEMPLATE  = os.path.join(ROOT, 'final', 'SRS_ManageBacklog.docx')
OUT       = os.path.join(ROOT, 'final', 'SRS_ManageWorkflow.docx')
TMP_LOGO  = os.path.join(ROOT, '_cover_logo.png')

# extract the cover logo (word/media/image1.png) from the template
with zipfile.ZipFile(TEMPLATE) as z:
    with open(TMP_LOGO, 'wb') as f:
        f.write(z.read('word/media/image1.png'))

doc = Document(TEMPLATE)

# wipe the template body, keep the trailing sectPr (page setup)
body = doc.element.body
for child in list(body):
    if child.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(child)

CONTENT_W     = Inches(6.3)
HEADING_COLOR = RGBColor(0x0F, 0x47, 0x61)
CAPTION_GREY  = RGBColor(0x6B, 0x77, 0x8C)
CENTER        = WD_ALIGN_PARAGRAPH.CENTER

# ── helpers ──────────────────────────────────────────────────────────────────
def cover_line(text, size, italic=True, black=False, before=None, after=None):
    p = doc.add_paragraph(); p.alignment = CENTER
    if before is not None: p.paragraph_format.space_before = Emu(before)
    if after  is not None: p.paragraph_format.space_after  = Emu(after)
    r = p.add_run(text)
    r.italic = italic
    r.font.name = 'Aptos Display'
    r.font.size = Pt(size)
    if black:
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def h(level, text):
    return doc.add_heading(text, level=level)

def normal(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold
    return p

def note(text):
    """Italic grey note paragraph (explains the no-image figure convention)."""
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True
    r.font.size = Pt(9); r.font.color.rgb = CAPTION_GREY
    return p

def bullet(text):
    return doc.add_paragraph('•  ' + text)

def figure_caption(caption):
    """No-image figure: the italic grey caption only (image to be placed)."""
    cap = doc.add_paragraph(); cap.alignment = CENTER
    cr = cap.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(9)
    cr.font.color.rgb = CAPTION_GREY

def set_cell_bold(cell):
    for para in cell.paragraphs:
        for r in para.runs:
            r.bold = True

def actors_table(rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for r_idx, (a, b) in enumerate(rows):
        c0, c1 = table.rows[r_idx].cells
        c0.text, c1.text = a, b
        c0.width = Emu(2971800); c1.width = Emu(2971800)
    set_cell_bold(table.rows[0].cells[0]); set_cell_bold(table.rows[0].cells[1])
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('bottom', 'insideH'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0');    el.set(qn('w:color'), 'D0D7DE')
        borders.append(el)
    tblPr.append(borders)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), '9360'); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)
    return table

# ═════════════════════════════════════════════════════════════════════════════
# COVER
# ═════════════════════════════════════════════════════════════════════════════
img_p = doc.add_paragraph(); img_p.alignment = CENTER
img_p.paragraph_format.space_before = Emu(1524000)
img_p.add_run().add_picture(TMP_LOGO, width=Inches(3.12))

cover_line('Technical Document', 24, black=True, before=381000, after=76200)
cover_line('Manage Workflow — SRS', 18, black=True, after=254000)
cover_line('SOFTWARE REQUIREMENTS SPECIFICATION', 11, after=38100)
cover_line('Manage Workflow · Workflow Designer', 11)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ═════════════════════════════════════════════════════════════════════════════
# BODY
# ═════════════════════════════════════════════════════════════════════════════
h(1, 'Software Requirements Specification (SRS) — Manage Workflow')
normal('Feature: Manage Workflow (Workflow designer) — Apex controller: '
       'force-app/main/default/classes/controller/manageWorkflow/ManageWorkflowPageController.cls '
       '— LWC: force-app/main/default/lwc/manageWorkflow/ — Date: 2026-06-22',
       bold=True)
note('Image references in this document use the figure-caption convention of the SRS series. '
     'This document ships without screenshots, so each figure is a one-line description of the '
     'image that should be placed there (no embedded image).')

# ── 1. Purpose & Scope ───────────────────────────────────────────────────────
h(2, '1. Purpose & Scope')
normal('The Manage Workflow page is the workflow-design surface of the Jira Clone. For a selected '
       'Project it lets a user author and maintain the workflows that govern how tickets move between '
       'statuses. From this page a user can:')
bullet('See the list of workflows defined for the selected project and open one to edit, or create a new one.')
bullet('See the chosen workflow as a visual graph: one node per project status and one curved arrow per transition.')
bullet('Create new statuses (nodes) for the project.')
bullet('Create a transition between two statuses by clicking the source then the target node.')
bullet('Inspect a transition (name, from/to status, record status, created date) and activate or delete it.')
bullet('Attach field-validation rules to a transition (the rules a ticket must satisfy to take that transition).')
bullet('Publish the draft by updating the workflow, which activates every pending transition at once.')
normal('The project is selected once and persisted in localStorage(’projectId’); if absent the page '
       'shows the Choose Project splash. The chosen workflow is persisted in localStorage(’workflowId’); '
       'the three entry states are: (1) no project → Choose Project, (2) project but no workflow → '
       'workflow list, (3) workflow selected → the visual editor.')
figure_caption('Figure 1 (image to be placed) — The workflow list step: the “Workflows” heading, a '
               '“Create New Workflow” button, and one card per workflow showing its name and record '
               'status with an Edit button.')

# ── 2. Actors ────────────────────────────────────────────────────────────────
h(2, '2. Actors')
actors_table([
    ('Actor', 'Description'),
    ('Project Member (User)',
     'The authenticated Salesforce user designing the project’s workflows — creates statuses and '
     'transitions, attaches validation rules, and activates or deletes transitions.'),
    ('System',
     'Apex services that persist statuses, workflows, transitions, and validation fields, enforce existence '
     'of referenced records, and toggle transition record status (pending → active, or soft-deleted).'),
])
doc.add_paragraph()

# ── 3. Functional Requirements ───────────────────────────────────────────────
h(2, '3. Functional Requirements')

h(3, '3.1 Project Selection and Workflow List')
bullet('FR-1 The system shall allow a user to select a project before accessing any workflow editor.')
bullet('FR-2 The system shall load and list the workflows defined for the selected project (name and record status).')
bullet('FR-3 When the project has no workflows, the system shall show an empty-state message inviting the user to create one.')
bullet('FR-4 The system shall allow a user to create a new workflow by name; on success the new workflow opens directly in the editor.')

h(3, '3.2 Workflow Editor Entry')
bullet('FR-5 The system shall allow a user to open an existing workflow in the visual editor by clicking Edit on its list row.')
bullet('FR-6 On entering the editor the system shall load the workflow configuration — its project statuses and its transitions — for the selected workflow.')
bullet('FR-7 The system shall allow a user to return from the editor to the workflow list.')

h(3, '3.3 Status (Node) Management')
bullet('FR-8 The system shall render each project status as a node on the editor canvas, labelled with the status name.')
bullet('FR-9 The system shall allow a user to create a new status by name; on success the status appears as a new node.')
bullet('FR-10 The canvas layout shall be responsive — the number of nodes per row adapts to the available width (4 / 3 / 2 columns for desktop / tablet / mobile).')

h(3, '3.4 Transition Creation')
bullet('FR-11 The system shall let a user define a transition by clicking a source status node (“From”) and then a different target status node (“To”); clicking the already-selected source again deselects it.')
bullet('FR-12 On selecting a second status the system shall open a Create Transition dialog showing From → To and asking for the transition name.')
bullet('FR-13 A newly created transition shall be recorded as pending and rendered as a dashed arrow until it is activated.')

h(3, '3.5 Transition Visualization')
bullet('FR-14 The system shall render active transitions as solid curved arrows and pending transitions as dashed curved arrows, each pointing from the source node to the target node.')
bullet('FR-15 The system shall show a legend distinguishing Active Transition from Pending Transition.')
bullet('FR-16 The system shall display each transition’s name as a label along its arrow.')

h(3, '3.6 Transition Detail and Lifecycle')
bullet('FR-17 The system shall let a user open a transition’s detail panel by clicking its arrow, showing the transition name, Id, From status, To status, record status, and created date.')
bullet('FR-18 The system shall let a user activate a pending transition (Activate is offered only while the transition is pending).')
bullet('FR-19 The system shall let a user delete a transition after a confirmation prompt; deletion is a soft delete and the arrow is removed from the canvas.')
bullet('FR-20 The system shall let a user close the detail panel.')

h(3, '3.7 Validation Rules')
bullet('FR-21 The system shall let a user add a field-validation rule to a transition by choosing a ticket field (API name) and a validation type.')
bullet('FR-22 The system shall let a user expand a transition to view its validation rules; the rules are loaded on expand.')

h(3, '3.8 Workflow Update / Activation')
bullet('FR-23 The system shall let a user update the workflow, which activates every pending transition of that workflow in a single action and returns to the workflow list.')

figure_caption('Figure 2 (image to be placed) — The visual editor: the toolbar (Back · Create Status · '
               'Update Workflow), the Active / Pending legend, status nodes laid out in a grid, and curved '
               'transition arrows (solid = active, dashed = pending) between them.')
figure_caption('Figure 3 (image to be placed) — Creating a transition: after clicking a source then a target '
               'node, the Create Transition dialog shows “From → To” and a Transition Name field.')
figure_caption('Figure 4 (image to be placed) — The transition detail panel: Name, Id, From / To status badges, '
               'Record Status badge, Created Date, with Activate (pending only) / Delete actions and the Add '
               'Validation Rule / Show Validation Details controls.')

# ── 4. Validation Rules ──────────────────────────────────────────────────────
h(2, '4. Validation Rules')
h(3, '4.1 Client-side (LWC validators)')
bullet('VR-1 A status name is required and must be 80 characters or fewer (validateStatusName).')
bullet('VR-2 A transition requires both a From and a To status, and the two must be different (validateTransition).')
bullet('VR-3 A transition name is required (validateTransitionName).')
bullet('VR-4 A validation rule requires both a ticket field and a validation type (validateTicketField, validateValidationType).')
bullet('VR-5 A transition must be selected before its validation details can be loaded (validateTransitionId).')
bullet('VR-6 Creating a workflow requires a non-blank name and a selected project; creating a transition requires a workflow to be open.')

h(3, '4.2 Server-side (controller input guards)')
bullet('VR-7 Every @AuraEnabled method blank-checks its required parameters and returns '
       'APIResponse(false, ’<field> is required’) before doing any work (e.g. workflowId, projectId, name, '
       'fromStatusId, toStatusId, workflowTransitionId, fieldName, type).')
bullet('VR-8 Referenced records must exist via DomainCorrectness.require* (requireProjectExists, '
       'requireWorkflowExists, requireStatusExists, requireTransitionExists) before any insert/update.')

# ── 5. Business Rules ────────────────────────────────────────────────────────
h(2, '5. Business Rules')
h(3, '5.1 Workflows & projects')
bullet('BR-1 A workflow belongs to exactly one project; the workflow list shows only that project’s workflows (loadWorkflowsByProject).')
bullet('BR-2 A status belongs to the project and is shared across that project’s workflows; the editor’s node set is the project’s statuses plus any status referenced by a transition.')

h(3, '5.2 Transition lifecycle')
bullet('BR-3 A transition is created in the pending record status and has no effect on ticket movement until it is active.')
bullet('BR-4 A transition becomes active either individually (Activate) or in bulk when the workflow is updated; Update Workflow activates all of that workflow’s pending transitions at once.')
bullet('BR-5 Deleting a transition is a soft delete (record status set to deleted): the arrow disappears but the underlying record is retained.')

h(3, '5.3 Validation rules')
bullet('BR-6 A validation rule (ValidateField__c) belongs to a single transition and records the ticket field plus the validation type that a ticket must satisfy to take that transition.')
bullet('BR-7 The selectable fields and types mirror the ValidateField__c restricted picklists (FieldName__c / Type__c), so the page only offers values the backend accepts.')

h(3, '5.4 Data hygiene')
bullet('BR-8 Workflow, status, transition, and validation reads exclude soft-deleted rows, per the project’s soql-exclude-deleted rule.')

# ── 6. Use Case Descriptions ─────────────────────────────────────────────────
h(2, '6. Use Case Descriptions')

h(3, 'UC-1 — Open a workflow for editing')
bullet('Actor: Project Member')
bullet('Pre-conditions: User is authenticated; a projectId exists in localStorage.')
bullet('Main flow:')
normal('1. The page loads and lists the project’s workflows (loadWorkflowsByProject). 2. User clicks Edit on '
       'a workflow. 3. The workflow Id is stored and getWorkflow loads its statuses and transitions. '
       '4. The visual editor renders the nodes and arrows.')
bullet('Alternative flows:')
bullet('A1 (no project): No projectId → Choose-Project splash; on projectchosen the project is stored and the list loads.')
bullet('A2 (no workflows): The list is empty → empty-state message; the user creates one (UC-2).')
bullet('A3 (load error): Apex returns success=false or throws → error toast.')

h(3, 'UC-2 — Create a workflow')
bullet('Actor: Project Member')
bullet('Pre-conditions: A project is selected; the workflow list is shown.')
bullet('Main flow:')
normal('1. User clicks Create New Workflow and enters a name. 2. createWorkflow(name, projectId) persists it. '
       '3. On success the new workflow Id is stored and the editor opens on the new (empty) workflow.')
bullet('Alternative flows:')
bullet('A1 (blank name): Required-field validation; no call.')
bullet('A2 (server error): Error toast; the modal stays open.')

h(3, 'UC-3 — Create a status')
bullet('Actor: Project Member')
bullet('Pre-conditions: The editor is open.')
bullet('Main flow:')
normal('1. User clicks Create Status and enters a name. 2. createStatus(name, projectId) persists it. '
       '3. The new status is added to the canvas as a node.')
bullet('Alternative flows:')
bullet('A1 (invalid name): Required / length validation (≤80 chars); no call.')

h(3, 'UC-4 — Create a transition')
bullet('Actor: Project Member')
bullet('Pre-conditions: The editor is open with at least two statuses.')
bullet('Main flow:')
normal('1. User clicks a source status node (From), then a different target node (To). 2. The Create Transition '
       'dialog opens showing From → To. 3. User enters a name and confirms. '
       '4. addWorkflowTransition(workflowId, name, fromStatusId, toStatusId) persists it. '
       '5. The transition appears as a pending (dashed) arrow.')
bullet('Alternative flows:')
bullet('A1 (same status / missing side): “From and To cannot be the same status” / “Both from and to '
       'statuses are required”; no call.')
bullet('A2 (blank name): “Transition name is required”; no call.')

h(3, 'UC-5 — View, activate, or delete a transition')
bullet('Actor: Project Member')
bullet('Pre-conditions: A transition arrow is visible.')
bullet('Main flow (any one of):')
bullet('View → click the arrow → the detail panel shows name, Id, from/to, record status, created date.')
bullet('Activate (pending only) → activateWorkflowTransition → the arrow becomes active (solid).')
bullet('Delete → confirm → deleteWorkflowTransition (soft delete) → the arrow is removed.')
bullet('Alternative flows:')
bullet('A1 (cancel delete): The confirmation is dismissed; nothing changes.')
bullet('A2 (server error): Error toast; the panel keeps its last good value.')

h(3, 'UC-6 — Add and view validation rules')
bullet('Actor: Project Member')
bullet('Pre-conditions: The transition detail panel is open.')
bullet('Main flow:')
normal('1. User clicks Add Validation Rule, chooses a Field API Name and a Validation Type, and confirms. '
       '2. addValidateField(workflowTransitionId, fieldName, type) persists the rule and it is folded onto the '
       'transition. 3. User clicks Show Validation Details to expand and list the transition’s rules '
       '(loadValidateFields).')
bullet('Alternative flows:')
bullet('A1 (missing field/type): Validation toast; no call.')
bullet('A2 (no transition selected on expand): “A transition must be selected before loading its validation details.”')

h(3, 'UC-7 — Update the workflow (activate pending)')
bullet('Actor: Project Member')
bullet('Pre-conditions: The editor is open and the workflow has one or more pending transitions.')
bullet('Main flow:')
normal('1. User clicks Update Workflow. 2. updateFullWorkflowTransition(workflowId) activates all pending '
       'transitions and reports the count. 3. The page returns to the workflow list.')
bullet('Alternative flows:')
bullet('A1 (no workflow id): “Workflow ID not found” error toast; no call.')
bullet('A2 (server error): Error toast; the editor stays open.')

# ── 7. Non-Functional Notes ──────────────────────────────────────────────────
h(2, '7. Non-Functional Notes')
bullet('Derived UI state: the LWC keeps one principal de-normalized state (workflowData = the workflow with '
       'its statuses and transitions) and derives every displayed value — the sorted status list, node '
       'positions, the SVG view box, the active/pending arrow geometry, the selected transition — via getters '
       'and pure utility functions, never parallel tracked fields.')
bullet('Responsive geometry: a ResizeObserver on the editor container recomputes the layout config when the '
       'width changes meaningfully (> 20px), so the graph re-flows between 4 / 3 / 2 columns without a server '
       'round-trip.')
bullet('Lazy validation reads: a transition’s validation rules are loaded only when the detail is expanded, '
       'behind a dedicated wire gate, so selecting a transition does not by itself fetch its rules.')
bullet('Type-safe option lists: the validation field/type comboboxes are sourced from constants that mirror the '
       'ValidateField__c restricted picklists, so the page emits exactly the API values the insert expects.')
bullet('Soft delete everywhere: transitions (and the other workflow records) are soft-deleted; reads filter them out.')

# ── save ─────────────────────────────────────────────────────────────────────
doc.save(OUT)
if os.path.exists(TMP_LOGO):
    os.remove(TMP_LOGO)
print('WROTE', OUT)
