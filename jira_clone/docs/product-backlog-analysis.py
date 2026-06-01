"""Generate product-backlog-analysis.xlsx from the 5 LWC modules analysis."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"c:\softreatail\project\app\jira_clone_v011\jira_clone\docs\product-backlog-analysis.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F3A8A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
GROUP_FILL  = PatternFill("solid", fgColor="E0E7FF")
GROUP_FONT  = Font(bold=True, color="1F3A8A", size=10)
WRAP        = Alignment(wrap_text=True, vertical="top")
THIN        = Side(style="thin", color="CBD5E1")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

def write_rows(ws, headers, rows, widths):
    ws.append(headers)
    style_header(ws, len(headers))
    for r in rows:
        # Group/separator rows: single-cell tuple with str
        if isinstance(r, tuple) and len(r) == 1:
            ws.append([r[0]] + [""] * (len(headers) - 1))
            row_idx = ws.max_row
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers))
            c = ws.cell(row=row_idx, column=1)
            c.fill = GROUP_FILL
            c.font = GROUP_FONT
            c.alignment = Alignment(vertical="center")
            ws.row_dimensions[row_idx].height = 20
            continue
        ws.append(list(r))
        row_idx = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = WRAP
            cell.border = BORDER
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ─── Summary sheet ────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Summary"
headers = ["Module", "Purpose", "Stories (count)", "State", "Key risks / gaps"]
rows = [
    ("manageBacklog",
     "Backlog + sprint planning board (create/edit sprints, ticket CRUD, drag-and-drop, peek panel).",
     33, "Production-ready (richest module)",
     "Ticket type is set-once; story-point edit not visible in this file (likely in c-ao-ticket-item)."),
    ("manageTicketTracking",
     "Active-sprint Kanban board with workflow-constrained drag between status lanes.",
     11, "Mostly working; one broken UI path",
     "Legacy c-ticket-linked-to popup wired in template but its handlers are missing in JS — dead code."),
    ("manageWorkflow",
     "Workflow designer: project → workflow list → SVG visualizer with statuses, transitions, activation.",
     14, "Functional, with one UX inconsistency",
     "Delete-transition uses native browser confirm() (inconsistent with the app's modal pattern). No rename-transition UI."),
    ("manageProjectMember",
     "Project member admin: search users, add, edit role, remove.",
     8,  "Production-ready",
     "Does NOT follow the project-picker pattern (no localStorage / c-choose-project) — owns its own combobox."),
    ("manageTicketType",
     "Ticket-type admin per project (name, description, icon, workflow).",
     7,  "Production-ready, minor bug",
     "Edit modal omits Description (create accepts it) — looks like a missed field, not by design."),
]
write_rows(ws, headers, rows, widths=[24, 50, 14, 32, 60])

# ─── Per-module sheets ────────────────────────────────────────────────────────
SHEET_HEADERS = ["ID", "Group", "User-visible feature / story", "Code reference"]
SHEET_WIDTHS  = [8, 22, 70, 60]

def backlog_rows():
    return [
        ("Project entry",),
        ("BL-01", "Project entry",
         "On first load, if no projectId in localStorage, show project picker.",
         "manageBacklog.js:229-244"),

        ("Sprints",),
        ("BL-02", "Sprints", "Create sprint (duration, start date, goal).", "manageBacklog.js:792-810"),
        ("BL-03", "Sprints", "Edit sprint (same fields).",                    "manageBacklog.js:767-789"),
        ("BL-04", "Sprints", "Start sprint (confirm).",                       "manageBacklog.js:894-906"),
        ("BL-05", "Sprints", "Complete sprint (confirm).",                    "manageBacklog.js:879-892"),
        ("BL-06", "Sprints", "Delete sprint; tickets auto-move back to backlog.",
                                                                              "manageBacklog.js:862-877, 1276-1283"),
        ("BL-07", "Sprints", "Expand / collapse sprint card.",                "manageBacklog.js:814-825"),
        ("BL-08", "Sprints", "Sprint ticket pagination (prev/next).",         "manageBacklog.js:827-839"),
        ("BL-09", "Sprints", "Story-point progress bar (ended / total / %).", "manageBacklog.html:119-127"),

        ("Tickets",),
        ("BL-10", "Tickets", "Create ticket directly in backlog.",            "manageBacklog.js:726-739"),
        ("BL-11", "Tickets", "Create ticket directly inside a sprint.",       "manageBacklog.js:708-724"),
        ("BL-12", "Tickets", "Inline edit: summary.",                         "manageBacklog.js:480-494"),
        ("BL-13", "Tickets", "Inline edit: priority (with validator).",       "manageBacklog.js:497-513"),
        ("BL-14", "Tickets", "Inline edit: status (workflow-aware; toast on end-status; sprint SP recalc).",
                                                                              "manageBacklog.js:516-546"),
        ("BL-15", "Tickets", "Inline edit: assignee.",                        "manageBacklog.js:549-565"),
        ("BL-16", "Tickets", "Inline edit: epic.",                            "manageBacklog.js:568-584"),
        ("BL-17", "Tickets", "Inline create epic from ticket (and auto-assign).",
                                                                              "manageBacklog.js:587-610"),
        ("BL-18", "Tickets", "Single ticket delete.",                         "manageBacklog.js:462-477"),
        ("BL-19", "Tickets", "Multi-select + bulk delete (with confirm).",    "manageBacklog.js:445-457, 426-440"),
        ("BL-20", "Tickets", "Backlog pagination (prev/next).",               "manageBacklog.js:1239-1260"),

        ("Drag & drop",),
        ("BL-21", "Drag & drop", "Drag ticket backlog → sprint.",             "manageBacklog.js:1150-1168"),
        ("BL-22", "Drag & drop", "Drag ticket sprint → backlog.",             "manageBacklog.js:1170-1182"),
        ("BL-23", "Drag & drop", "Reorder ticket within container (drop on ticket or top zone).",
                                                                              "manageBacklog.js:1095-1134"),
        ("BL-24", "Drag & drop", "Sprint→sprint cross-move blocked by design (dropEffect='none').",
                                                                              "manageBacklog.js:1008-1013"),

        ("Subtasks",),
        ("BL-25", "Subtasks", "Create subtask from ticket row.",              "manageBacklog.js:615-629"),
        ("BL-26", "Subtasks", "Inline edit subtask summary.",                 "manageBacklog.js:632-643"),
        ("BL-27", "Subtasks", "Change subtask assignee.",                     "manageBacklog.js:646-659"),
        ("BL-28", "Subtasks", "Single subtask delete + bulk subtask delete.", "manageBacklog.js:662-687"),

        ("Ticket-view peek panel",),
        ("BL-29", "Peek panel", "Open / close peek panel from a ticket.",     "manageBacklog.js:258-266"),
        ("BL-30", "Peek panel", "Update summary / description / status from peek.",
                                                                              "manageBacklog.js:268-305"),
        ("BL-31", "Peek panel", "Linked-to: lazy load on expand, search, create link.",
                                                                              "manageBacklog.js:307-334"),
        ("BL-32", "Peek panel", "Subtasks: lazy load on expand + create from peek.",
                                                                              "manageBacklog.js:336-356"),

        ("UX / infra",),
        ("BL-33", "Infra", "Responsive ticket variant via matchMedia (row vs full-ticket-card).",
                                                                              "manageBacklog.js:224-227, 750"),
    ]

def tracking_rows():
    return [
        ("Project entry",),
        ("TT-01", "Project entry", "Project picker entry (localStorage or c-choose-project).",
                                                                              "manageTicketTracking.js:257-274"),
        ("Sprint header",),
        ("TT-02", "Sprint header", "Show active sprint name, dates, goal.",   "manageTicketTracking.html:26-43"),
        ("TT-03", "Sprint header", "Story-point progress (ended / total / %).",
                                                                              "manageTicketTracking.html:33-41"),
        ("TT-04", "Sprint header", "“No active sprint” banner when none in progress.",
                                                                              "manageTicketTracking.html:45-47"),
        ("Kanban board",),
        ("TT-05", "Kanban", "Derive columns from project statuses; tickets fill lanes.",
                                                                              "manageTicketTracking.js:242-244"),
        ("TT-06", "Kanban", "Drag ticket between status columns — only along valid workflow transitions.",
                                                                              "manageTicketTracking.js:338-366"),
        ("TT-07", "Kanban", "Sprint TotalEnded SP recalculated when ticket lands in an end-status.",
                                                                              "manageTicketTracking.js:317-325"),
        ("Ticket peek panel",),
        ("TT-08", "Peek panel", "Open ticket-view side panel.",               "manageTicketTracking.js:398-400"),
        ("TT-09", "Peek panel", "Edit summary / description / status from peek.",
                                                                              "manageTicketTracking.js:404-433"),
        ("TT-10", "Peek panel", "Linked-to: lazy load on expand, search by term, create link.",
                                                                              "manageTicketTracking.js:437-461"),
        ("TT-11", "Peek panel", "Subtasks: lazy load on expand + create.",    "manageTicketTracking.js:467-485"),
    ]

def workflow_rows():
    return [
        ("Project & workflow entry",),
        ("WF-01", "Entry", "Three-step entry: project → workflow list → visualizer.",
                                                                              "manageWorkflow.js:71-82"),
        ("WF-02", "Entry", "Load workflows for current project.",             "manageWorkflow.js:98-120"),
        ("WF-03", "Entry", "Create new workflow (modal, name required).",     "manageWorkflow.js:184-230"),
        ("WF-04", "Entry", "Click ‘Edit’ to enter visualizer for a workflow.","manageWorkflow.js:122-132"),
        ("Visualizer",),
        ("WF-05", "Visualizer", "Load full workflow (statuses + transitions).",
                                                                              "manageWorkflow.js:257-271"),
        ("WF-06", "Visualizer", "Render SVG canvas: statuses, active and pending transitions, legend.",
                                                                              "manageWorkflow.html:158-285"),
        ("WF-07", "Visualizer", "ResizeObserver-driven responsive geometry.","manageWorkflow.js:232-250"),
        ("Statuses",),
        ("WF-08", "Statuses", "Create status (modal, validated).",            "manageWorkflow.js:408-433"),
        ("WF-09", "Statuses", "Click two statuses → opens Create-transition modal.",
                                                                              "manageWorkflow.js:367-381"),
        ("Transitions",),
        ("WF-10", "Transitions", "Create new transition, saved as ‘pending’ (not active).",
                                                                              "manageWorkflow.js:486-491"),
        ("WF-11", "Transitions", "Click transition line → side panel with full details.",
                                                                              "manageWorkflow.js:556-586"),
        ("WF-12", "Transitions", "Activate a single pending transition.",     "manageWorkflow.js:588-612"),
        ("WF-13", "Transitions", "Delete transition (native browser confirm).",
                                                                              "manageWorkflow.js:614-642"),
        ("WF-14", "Transitions", "‘Update Workflow’ bulk-activates all pending transitions.",
                                                                              "manageWorkflow.js:157-181"),
    ]

def member_rows():
    return [
        ("PM-01", "Projects",  "Load all projects, auto-select the first.",  "manageProjectMember.js:46-60"),
        ("PM-02", "Projects",  "Switch active project via combobox → reload members.",
                                                                              "manageProjectMember.js:76-84"),
        ("PM-03", "Members",   "List members (name + role) for the selected project.",
                                                                              "manageProjectMember.html:66-95"),
        ("PM-04", "Search",    "Debounced (180 ms) user search w/ dropdown of candidates.",
                                                                              "manageProjectMember.js:115-131"),
        ("PM-05", "Members",   "Add member from search hit (with already-member guard).",
                                                                              "manageProjectMember.js:173-197"),
        ("PM-06", "Members",   "Edit member role via modal (role allow-list validation).",
                                                                              "manageProjectMember.js:201-236"),
        ("PM-07", "Members",   "Delete member with confirm modal.",           "manageProjectMember.js:259-283"),
        ("PM-08", "UX",        "Inline toast (success / error / default), auto-dismiss 2.6 s.",
                                                                              "manageProjectMember.js:301-313"),
    ]

def tickettype_rows():
    return [
        ("TY-01", "Project entry", "Project picker entry (localStorage).",    "manageTicketType.js:45-62"),
        ("TY-02", "List",          "Load ticket types + workflows for project.",
                                                                              "manageTicketType.js:65-92"),
        ("TY-03", "List",          "List ticket types: icon, description, workflow name.",
                                                                              "manageTicketType.html:34-75"),
        ("TY-04", "Create",        "Create ticket type (name, description, icon URL, workflow).",
                                                                              "manageTicketType.js:138-167"),
        ("TY-05", "Edit",          "Edit ticket type (name, icon URL, workflow) — DOES NOT include description.",
                                                                              "manageTicketType.js:196-225"),
        ("TY-06", "Delete",        "Delete ticket type with confirm modal.",  "manageTicketType.js:242-274"),
        ("TY-07", "UX",            "Toast feedback (auto-dismiss 2.6 s).",    "manageTicketType.js:282-294"),
    ]

write_rows(wb.create_sheet("manageBacklog"),         SHEET_HEADERS, backlog_rows(),    SHEET_WIDTHS)
write_rows(wb.create_sheet("manageTicketTracking"),  SHEET_HEADERS, tracking_rows(),   SHEET_WIDTHS)
write_rows(wb.create_sheet("manageWorkflow"),        SHEET_HEADERS, workflow_rows(),   SHEET_WIDTHS)
write_rows(wb.create_sheet("manageProjectMember"),   SHEET_HEADERS, member_rows(),     SHEET_WIDTHS)
write_rows(wb.create_sheet("manageTicketType"),      SHEET_HEADERS, tickettype_rows(), SHEET_WIDTHS)

# ─── Gaps / next-iteration backlog ────────────────────────────────────────────
gaps_ws = wb.create_sheet("Gaps & next-iteration")
gap_headers = ["ID", "Module", "Severity", "Gap / candidate story", "Where / evidence"]
gap_rows = [
    ("G-01", "manageTicketTracking", "High",
     "Linked-to popup is dead code: template wires onopenlinkedto / onexpandlinkedto / oncloselinkedto / oncreateticketlink but the JS class no longer defines those handlers. Either delete the markup or restore the handlers.",
     "manageTicketTracking.html:51, 72-83"),
    ("G-02", "manageTicketType",     "Medium",
     "Edit modal is missing the Description field even though Create accepts it. Looks like a bug, not by design.",
     "manageTicketType.html:140-181 vs js:138-167"),
    ("G-03", "manageWorkflow",       "Medium",
     "Delete-transition uses native browser confirm() — inconsistent with the in-app modal pattern used everywhere else.",
     "manageWorkflow.js:621"),
    ("G-04", "manageWorkflow",       "Medium",
     "No ‘rename transition’ UI — only activate or delete are exposed in the detail panel, despite controller support.",
     "manageWorkflow.js:556-642"),
    ("G-05", "manageProjectMember",  "Low",
     "Doesn’t follow the project-picker pattern (no localStorage / c-choose-project); owns its own combobox. Consider aligning for consistency.",
     "manageProjectMember.js:42-84"),
    ("G-06", "manageBacklog",        "Low",
     "Ticket type (Ticket_Type__c) is set-once at create — not editable from the ticket row.",
     "manageBacklog.js (no handleTicketTypeUpdate)"),
    ("G-07", "manageBacklog",        "Low",
     "Story-point editor for an existing ticket isn’t visible at this level. Sprint SP totals depend on it — confirm it’s covered by c-ao-ticket-item or treat as a gap.",
     "manageBacklog.js (no handleTicketStoryPointUpdate)"),
]
write_rows(gaps_ws, gap_headers, gap_rows, widths=[8, 22, 12, 70, 50])

wb.save(OUT)
print(f"WROTE {OUT}")
