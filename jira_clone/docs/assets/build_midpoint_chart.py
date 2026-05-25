"""Build a chart from midpoint_sequence.xlsx showing how the gap between
the midpoint and the lower bound collapses toward zero in float64 — used as
visual evidence in the Technical Design Document."""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = Path(r"c:/softreatail/project/app/jira_clone_v011/jira_clone")
XLSX = ROOT / "midpoint_sequence.xlsx"
OUT = ROOT / "docs" / "assets" / "midpoint_collapse.png"

wb = load_workbook(XLSX, data_only=True)
ws = wb["Midpoint Sequence"]

steps, gaps = [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    step, left, right, mid = row
    if step is None:
        continue
    gap = float(mid) - float(left)
    steps.append(int(step))
    gaps.append(gap if gap > 0 else 1e-20)  # floor to keep log scale plotting

fig, ax = plt.subplots(figsize=(9, 4.5))
ax.semilogy(steps, gaps, marker="o", markersize=3.5, linewidth=1.4, color="#0b5394")
ax.set_xlabel("Step (number of midpoints taken between 1 and 2)")
ax.set_ylabel("Gap = midpoint − 1   (log scale, float64)")
ax.set_title("Floating-point midpoint sequence collapses to zero after ~52 steps")
ax.grid(True, which="both", linestyle=":", alpha=0.5)

# annotate the collision step
collision_step = None
for s, g in zip(steps, gaps):
    if g <= 1e-20:
        collision_step = s
        break
if collision_step:
    ax.axvline(collision_step, color="#cc0000", linewidth=1, linestyle="--")
    ax.annotate(
        f"Collision (gap = 0)\nstep {collision_step}",
        xy=(collision_step, 1e-16),
        xytext=(collision_step - 22, 1e-12),
        fontsize=9,
        color="#cc0000",
        arrowprops=dict(arrowstyle="->", color="#cc0000", lw=0.8),
    )

fig.tight_layout()
fig.savefig(OUT, dpi=160)
print(f"wrote {OUT}")
